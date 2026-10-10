import csv
import io

from django.http import HttpResponse
from . models import Hotel


COLUMNS = ['ID', 'Название', 'Город', 'Адрес', 'Цена за ночь', 'Рейтинг']

def _hotel_rows():
    for hotel in Hotel.objects.filter(is_active=True).order_by('name'):
        yield [hotel.pk, hotel.name, hotel.city, hotel.address, hotel.price_night, hotel.rating]


def export_hotels_csv(request):
    response = HttpResponse(content_type='text/csv; charset=utf-8')   
    response['Content-Disposition'] = 'attachment; filename="hotels.csv"'
    response.write('\ufeff')

    writer = csv.writer(response, delimiter=';')
    writer.writerow(COLUMNS)
    for row in _hotel_rows():
        writer.writerow(row)
    return response


def export_hotels_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Отели'

    ws.append(COLUMNS)
    for cell in ws[1]:
        ws.font = Font(bold=True)

    for row in _hotel_rows():
        ws.append(row)

    for col_cells in ws.columns:
        lenght = max(len(str(c.value)) for c in col_cells if c.value is not None)
        ws.column_dimensions[col_cells[0].column_letter].width = lenght + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="hotels.xlsx"'
    return response


def export_hotels_pdf(request):
    
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)

    styles = getSampleStyleSheet()
    elements = [Paragraph('HotelHub — каталог отелей', styles['Title']), Spacer(1, 12)]

    data = [COLUMNS] + [[str(v) for v in row] for row in _hotel_rows()]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#161B2E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F4F6FC')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="hotels.pdf"'
    return response


def export_hotels_docx(request):
    """
    Экспорт в Word (.docx) через python-docx — редактируемый документ,
    удобный, например, персоналу для распечатки или правки офлайн.
    """
    from docx import Document

    document = Document()
    document.add_heading('HotelHub — каталог отелей', level=1)

    table = document.add_table(rows=1, cols=len(COLUMNS))
    table.style = 'Light Grid Accent 1'
    header_cells = table.rows[0].cells
    for i, col_name in enumerate(COLUMNS):
        header_cells[i].text = col_name

    for row in _hotel_rows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = 'attachment; filename="hotels.docx"'
    return response